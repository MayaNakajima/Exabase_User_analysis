"""STEP 16: Excel集計ブック（全期間・月次）→ output/エクサベースAI_利用分析_全期間.xlsx"""
import os, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from config import DATA, OUTPUT
os.chdir(DATA)   # 相対パスの read/write は data/ に解決

import json
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

agg=json.load(open('aggregates_all.json'))
rows=json.load(open('user_all.json'))
meta=json.load(open('meta_all.json'))
CATN=meta['cat_names']; GL=meta['grp_label']
F='Arial'
HDR=Font(name=F,bold=True,color='FFFFFF',size=10); HF=PatternFill('solid',fgColor='1F4E78')
TIT=Font(name=F,bold=True,size=14,color='1F4E78'); SUB=Font(name=F,bold=True,size=11,color='1F4E78')
NM=Font(name=F,size=10); BLD=Font(name=F,bold=True,size=10)
thin=Side(style='thin',color='D9D9D9'); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
CEN=Alignment('center','center'); LEFT=Alignment('left','center',wrap_text=True); L2=Alignment('left','center')
QF={'Q1':'C6E0B4','Q2':'FFE699','Q3':'BDD7EE','Q4':'F8CBAD'}
GF={'core':'C6E0B4','light':'BDD7EE','user':'FCE4D6','former':'EDEDED','none':'FFFFFF'}
wb=Workbook()

def hdrrow(ws,row,heads,fromcol=1):
    for i,h in enumerate(heads,fromcol):
        c=ws.cell(row=row,column=i,value=h); c.font=HDR; c.fill=HF; c.alignment=CEN; c.border=BD
def put(ws,r,c,v,font=NM,al=L2,fill=None,bd=True,fmt=None):
    cell=ws.cell(row=r,column=c,value=v); cell.font=font; cell.alignment=al
    if fill: cell.fill=PatternFill('solid',fgColor=fill)
    if bd: cell.border=BD
    if fmt: cell.number_format=fmt
    return cell

t=agg['totals']
# ---------- サマリー ----------
ws=wb.active; ws.title='サマリー'; ws.sheet_view.showGridLines=False
put(ws,1,1,'エクサベースAI 利用実態分析（全期間・月次）',TIT,L2,bd=False)
put(ws,2,1,f"対象: {agg['months'][0]}〜{agg['months'][-1]}（{t['months']}ヶ月） ／ 利用ユーザー {t['users']}名 ／ 会話 {t['convs']:,}件 ／ メッセージ {t['msgs']:,}件",NM,L2,bd=False)
r=4; put(ws,r,1,'■ 全体サマリー',SUB,L2,bd=False); r+=1
tt=t['eco']+t['std']+t['pre']
for k,v in [('メンバー総数',f"{t['members']}名（稼働 {t['active_members']}／非稼働 {t['members']-t['active_members']}）"),
            ('利用実績のあるユーザー',f"{t['users']}名"),
            ('総会話数',f"{t['convs']:,}件"),('総メッセージ数',f"{t['msgs']:,}件"),
            ('Economy 発話',f"{t['eco']:,}（{t['eco']/tt*100:.0f}%）"),
            ('Standard 発話',f"{t['std']:,}（{t['std']/tt*100:.0f}%）"),
            ('Premium 発話',f"{t['pre']:,}（{t['pre']/tt*100:.1f}%）"),
            ('過剰スペック候補（Claude併用4名除く）',f"{agg['overspec_count']:,}件"),
            ('月次アクティブユーザー',f"{agg['monthly'][0]['active_users']}名（{agg['months'][0]}）→ {agg['monthly'][-1]['active_users']}名（{agg['months'][-1]}）")]:
    put(ws,r,1,k,NM); put(ws,r,2,v,BLD); r+=1
r+=1; put(ws,r,1,'■ 主要インサイト',SUB,L2,bd=False); r+=1
gq={g['group']:g for g in agg['groupquad']}
ins=[
 f"① 導入は着実に拡大。月次アクティブは {agg['monthly'][0]['active_users']}→{agg['monthly'][-1]['active_users']}名、月次会話は {agg['monthly'][0]['convs']}→{agg['monthly'][-1]['convs']}件。累計利用者は {agg['monthly'][-1]['cum_users']}名。社内活動の効果測定はこの月次推移で継続的に追える。",
 f"② 管理グループと利用実態（4象限）は概ね整合。高頻度・高リテラシー(Q1)比率は コア {gq.get('core',{}).get('Q1',0)}/{gq.get('core',{}).get('total',0)} ＞ ライト {gq.get('light',{}).get('Q1',0)}/{gq.get('light',{}).get('total',0)} ＞ ユーザー {gq.get('user',{}).get('Q1',0)}/{gq.get('user',{}).get('total',0)}。",
 f"③ 一方でズレも存在。ライト層の低活用(Q4)が {gq.get('light',{}).get('Q4',0)}名（離脱・定着支援の対象）、ユーザー層の高活用(Q1)が {gq.get('user',{}).get('Q1',0)}名（コア候補・横展開の担い手）。",
 "④ リテラシーとStandard/Premium(従量課金)選択はほぼ無相関(r=0.10)。上位モデルは必要性ではなく習熟度と無関係に選ばれており、モデル推奨ガイドで最適化余地が大きい。",
 f"⑤ 過剰スペック候補は全期間で {agg['overspec_count']:,}件。単純用途はEconomy既定とする運用で継続的なコスト削減が可能。",
 "⑥ 公平性のため、利用頻度は「稼働月あたりの発話数」で正規化。入退社・利用開始時期の差を排除して比較している。",
]
for s in ins:
    put(ws,r,1,s,NM,LEFT,bd=False); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=7); ws.row_dimensions[r].height=46; r+=1
r+=1; put(ws,r,1,'■ 分析手法・注記',SUB,L2,bd=False); r+=1
notes=[
 "リテラシー複合スコア(0-100)=会話の深さ18%+プロンプト具体性18%+高度ツール活用率20%+用途多様性16%+構造化プロンプト率18%+ツール種類数10%（各指標を利用者内百分位に変換し加重平均）。",
 "利用頻度=稼働月あたりのユーザー発話数（在籍・利用期間の差を正規化）。4象限は頻度・リテラシーの中央値で分割。",
 "グループ（コア/ライト/ユーザー）はメンバー一覧のサブグループ権限に基づく。メンバー表に無い利用者は「退職/不明」。",
 "内容分類は全29,548会話の初回発話をAIが12カテゴリに分類。過剰スペック候補はStandard/Premium×単純用途×往復2以下×高度ツール無し（Claude利用許可4名は除外）。",
 "Claude併用4名（a-miyairi / kazunori.matsuyama / kentaro.takeishi / maya.nakajima）は高度用途をClaudeにオフロードするため、ExaBaseの上位モデル利用・リテラシー解釈に注記。",
]
for s in notes:
    put(ws,r,1,s,Font(name=F,size=9,color='595959'),LEFT,bd=False); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=7); ws.row_dimensions[r].height=40; r+=1
ws.column_dimensions['A'].width=40
for col in 'BCDEFG': ws.column_dimensions[col].width=15

# ---------- 月次推移 ----------
ws=wb.create_sheet('月次推移'); ws.sheet_view.showGridLines=False
put(ws,1,1,'月次推移（利用量・ユーザー・モデル階層・過剰スペック）',SUB,L2,bd=False)
heads=['月','アクティブ\nユーザー','新規\nユーザー','累計\nユーザー','会話数','メッセージ数','ユーザー\nあたりmsg','Economy\n発話','Standard\n発話','Premium\n発話','Std/Prem\n比率%','高度ツール\n活用率%','過剰スペック\n候補']
hdrrow(ws,3,heads)
rr=4
for m in agg['monthly']:
    vals=[m['month'],m['active_users'],m['new_users'],m['cum_users'],m['convs'],m['msgs'],m['msgs_per_user'],
          m['eco'],m['std'],m['pre'],m['sp_share'],m['adv_rate'],m['overspec']]
    for i,v in enumerate(vals,1): put(ws,rr,i,v,NM,CEN if i>1 else L2)
    rr+=1
ws.freeze_panes='B4'; ws.row_dimensions[3].height=30
ws.column_dimensions['A'].width=10
for i in range(2,14): ws.column_dimensions[get_column_letter(i)].width=11

# ---------- 月次×グループ ----------
ws=wb.create_sheet('月次×グループ'); ws.sheet_view.showGridLines=False
put(ws,1,1,'月次 グループ別アクティブユーザー数・発話数',SUB,L2,bd=False)
heads=['月','コア\n人数','ライト\n人数','ユーザー\n人数','退職/不明\n人数','コア\n発話','ライト\n発話','ユーザー\n発話','退職/不明\n発話']
hdrrow(ws,3,heads)
rr=4
for m in agg['monthly_group']:
    vals=[m['month'],m['core_users'],m['light_users'],m['user_users'],m['former_users'],
          m['core_msgs'],m['light_msgs'],m['user_msgs'],m['former_msgs']]
    for i,v in enumerate(vals,1): put(ws,rr,i,v,NM,CEN if i>1 else L2)
    rr+=1
ws.freeze_panes='B4'; ws.row_dimensions[3].height=30
ws.column_dimensions['A'].width=10
for i in range(2,10): ws.column_dimensions[get_column_letter(i)].width=11

# ---------- ユーザー別（全期間） ----------
ws=wb.create_sheet('ユーザー別_全期間'); ws.sheet_view.showGridLines=False
heads=['ユーザー','氏名','グループ','稼働','Claude\n併用','初回月','最終月','稼働\n月数','会話数','発話数','月あたり\n発話','リテラシー','4象限','平均\n往復','高度ツール%','構造化%','用途数','Std/Prem\n発話%','過剰\n候補']
hdrrow(ws,1,heads)
rr=2
for u in sorted(rows,key=lambda x:(-x['literacy'],-x['msgs_per_month'])):
    vals=[u['user'],u['name'],GL.get(u['group'],u['group']),'稼働' if u['active'] else '非稼働','●' if u['claude'] else '',
          u['first_month'],u['last_month'],u['active_months'],u['convs'],u['umsgs'],round(u['msgs_per_month'],1),
          u['literacy'],u['quadrant'],round(u['avg_turns'],1),round(u['adv_rate']*100,1),round(u['struct_rate']*100,1),
          u['cat_div'],round(u['sp_share']*100),u['overspec']]
    for i,v in enumerate(vals,1):
        cell=put(ws,rr,i,v,NM,L2 if i in(1,2) else CEN)
        if i==3: cell.fill=PatternFill('solid',fgColor=GF.get(u['group'],'FFFFFF'))
        if i==13: cell.fill=PatternFill('solid',fgColor=QF[u['quadrant']])
        if i==12: cell.font=BLD
    rr+=1
ws.freeze_panes='C2'; ws.auto_filter.ref=f'A1:{get_column_letter(len(heads))}{rr-1}'
W=[24,14,13,8,7,9,9,7,8,9,9,10,7,7,10,9,7,10,7]
for i,w in enumerate(W,1): ws.column_dimensions[get_column_letter(i)].width=w
ws.row_dimensions[1].height=30

# ---------- グループ×象限 ----------
ws=wb.create_sheet('グループ×4象限'); ws.sheet_view.showGridLines=False
put(ws,1,1,'管理グループ × 利用実態（4象限）クロス',SUB,L2,bd=False)
put(ws,2,1,'※ 対角に近いほど管理グループと実態が整合。オフ対角が施策対象（ライト×Q4=定着支援 / ユーザー×Q1=コア候補）',Font(name=F,size=9,color='595959'),L2,bd=False)
heads=['グループ','Q1 高頻度・高ﾘﾃﾗｼｰ','Q2 高頻度・低ﾘﾃﾗｼｰ','Q3 低頻度・高ﾘﾃﾗｼｰ','Q4 低頻度・低ﾘﾃﾗｼｰ','合計','Q1比率%']
hdrrow(ws,4,heads)
rr=5
for g in agg['groupquad']:
    q1r=round(g['Q1']/g['total']*100) if g['total'] else 0
    vals=[g['label'],g['Q1'],g['Q2'],g['Q3'],g['Q4'],g['total'],q1r]
    for i,v in enumerate(vals,1):
        cell=put(ws,rr,i,v,NM,L2 if i==1 else CEN)
        if i==1: cell.fill=PatternFill('solid',fgColor=GF.get(g['group'],'FFFFFF')); cell.font=BLD
    rr+=1
ws.column_dimensions['A'].width=16
for c in 'BCDE': ws.column_dimensions[c].width=20
ws.column_dimensions['F'].width=8; ws.column_dimensions['G'].width=10
ws.row_dimensions[4].height=28

# ---------- 4象限サマリー ----------
ws=wb.create_sheet('4象限サマリー'); ws.sheet_view.showGridLines=False
put(ws,1,1,'4象限サマリーと推奨アクション',SUB,L2,bd=False)
heads=['象限','人数','会話数','発話数','平均リテラシー','高度ツール%','過剰候補','推奨アクション']
hdrrow(ws,3,heads)
ACT={'Q1':'横展開の担い手。コア認定・ベストプラクティス共有・高度機能(Deep Research/RAG)への引き上げ。',
'Q2':'量は使うが質に課題。プロンプト研修・テンプレート配布で費用対効果を改善。研修効果が最大の層。',
'Q3':'スキルはあるが機会が少ない。業務プロセスへの組み込み・活用シーン提示で頻度を引き上げ。',
'Q4':'未活用層。オンボーディング・成功体験設計。シャドウAI利用のリスク調査対象。'}
rr=4
for q in agg['quads']:
    vals=[f"{q['q']} {q['name']}",q['users'],q['convs'],q['msgs'],q['avg_lit'],q['avg_adv'],q['overspec'],ACT[q['q']]]
    for i,v in enumerate(vals,1):
        cell=put(ws,rr,i,v,NM,L2 if i in(1,8) else CEN)
        if i==1: cell.fill=PatternFill('solid',fgColor=QF[q['q']]); cell.font=BLD
        if i==8: cell.alignment=LEFT
    ws.row_dimensions[rr].height=48; rr+=1
ws.column_dimensions['A'].width=22
for c in 'BCDEFG': ws.column_dimensions[c].width=12
ws.column_dimensions['H'].width=62

# ---------- チャットパートナー別 ----------
ws=wb.create_sheet('チャットパートナー別'); ws.sheet_view.showGridLines=False
put(ws,1,1,'チャットパートナー（モデル）別 利用状況（全期間）',SUB,L2,bd=False)
hdrrow(ws,3,['チャットパートナー','階層','会話数','メッセージ数'])
TFILL={'Economy':'E2EFDA','Standard':'FFF2CC','Premium':'FCE4D6','Unknown':'FFFFFF'}
rr=4
for p in agg['partners']:
    put(ws,rr,1,p['partner'],NM,L2); put(ws,rr,2,p['tier'],NM,CEN,fill=TFILL.get(p['tier'],'FFFFFF'))
    put(ws,rr,3,p['convs'],NM,CEN); put(ws,rr,4,p['msgs'],NM,CEN); rr+=1
ws.column_dimensions['A'].width=46
for c in 'BCD': ws.column_dimensions[c].width=15

# ---------- 利用ツール別 ----------
ws=wb.create_sheet('利用ツール別'); ws.sheet_view.showGridLines=False
put(ws,1,1,'利用ツール別 利用会話数（全期間）',SUB,L2,bd=False)
hdrrow(ws,3,['利用ツール','利用会話数'])
rr=4
for tl in agg['tools']:
    put(ws,rr,1,tl['tool'],NM,L2); put(ws,rr,2,tl['convs'],NM,CEN); rr+=1
ws.column_dimensions['A'].width=42; ws.column_dimensions['B'].width=14

# ---------- 分類別 ----------
ws=wb.create_sheet('チャット内容分類別'); ws.sheet_view.showGridLines=False
put(ws,1,1,'チャット内容分類別 利用状況と階層内訳（全期間）',SUB,L2,bd=False)
hdrrow(ws,3,['コード','業務カテゴリ','会話数','構成比%','Economy','Standard','Premium'])
tot=sum(c['convs'] for c in agg['cats']); rr=4
for c in agg['cats']:
    vals=[c['code'],c['name'],c['convs'],round(c['convs']/tot*100,1),c['economy'],c['standard'],c['premium']]
    for i,v in enumerate(vals,1): put(ws,rr,i,v,NM,L2 if i==2 else CEN)
    rr+=1
ws.column_dimensions['A'].width=8; ws.column_dimensions['B'].width=30
for c in 'CDEFG': ws.column_dimensions[c].width=12

# ---------- Claude併用者 ----------
ws=wb.create_sheet('Claude併用者'); ws.sheet_view.showGridLines=False
put(ws,1,1,'Claude利用許可者（高度用途はClaudeにオフロード）',SUB,L2,bd=False)
put(ws,2,1,'※ ExaBaseの上位モデル利用が少なくても能力不足ではない。過剰スペック集計からは除外済み。',Font(name=F,size=9,color='595959'),L2,bd=False)
hdrrow(ws,4,['ユーザー','氏名','グループ','会話数','発話数','リテラシー','4象限','Std/Prem%'])
rr=5
for u in agg['claude_users']:
    vals=[u['user'],u['name'],GL.get(u['group'],u['group']),u['convs'],u['umsgs'],u['literacy'],u['quadrant'],u['sp_share']]
    for i,v in enumerate(vals,1): put(ws,rr,i,v,NM,L2 if i in(1,2) else CEN)
    rr+=1
ws.column_dimensions['A'].width=28; ws.column_dimensions['B'].width=14
for c in 'CDEFGH': ws.column_dimensions[c].width=11

# ---------- 過剰スペック候補 ----------
ws=wb.create_sheet('過剰スペック候補'); ws.sheet_view.showGridLines=False
put(ws,1,1,f"過剰スペック候補（Std/Prem×単純用途×浅い往復×高度ツール無し, Claude4名除く） 全{agg['overspec_count']:,}件",SUB,L2,bd=False)
put(ws,2,1,'※ 表示は先頭400件。全件はスクリプト再実行で出力可。ユーザー別の集計はユーザー別シートの「過剰候補」列を参照。',Font(name=F,size=9,color='595959'),L2,bd=False)
hdrrow(ws,4,['ユーザー','月','モデル','階層','用途','初回発話(抜粋)'])
rr=5
for o in agg['overspec']:
    vals=[o['user'],o['month'],o['model'],o['tier'],o['cat'],o['q'].replace('\n',' ')]
    for i,v in enumerate(vals,1): put(ws,rr,i,v,NM,L2 if i in(1,3,5,6) else CEN)
    rr+=1
ws.freeze_panes='A5'
ws.column_dimensions['A'].width=24; ws.column_dimensions['B'].width=9; ws.column_dimensions['C'].width=24
ws.column_dimensions['D'].width=10; ws.column_dimensions['E'].width=20; ws.column_dimensions['F'].width=64

wb.save(str(OUTPUT/'エクサベースAI_利用分析_全期間.xlsx'))
print('saved xlsx; sheets:', wb.sheetnames)
