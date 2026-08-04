"""STEP 6: Excel集計ブックを作成 → output/エクサベースAI_利用分析_202607.xlsx

サマリー/ユーザー別/パートナー別/ツール別/分類別/4象限/過剰スペック候補 の7シート。
生成後は scripts/recalc(LibreOffice) 等で数式検証すること(数式はごく僅か)。
"""
from config import DATA, OUTPUT, CAT_NAMES as _CN

import json
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

convs=json.load(open(DATA/'conversations.json'))
cls=json.load(open(DATA/'classification.json'))
rows=json.load(open(DATA/'user_rows.json'))
agg=json.load(open(DATA/'aggregates.json'))
meta=json.load(open(DATA/'meta.json'))
CAT=meta['cat_names']
for c in convs: c['cat']=cls.get(c['conv_id'],'C12')

FONT='Arial'
HDR=Font(name=FONT,bold=True,color='FFFFFF',size=10)
HDRFILL=PatternFill('solid',fgColor='1F4E78')
TITLE=Font(name=FONT,bold=True,size=14,color='1F4E78')
SUB=Font(name=FONT,bold=True,size=11,color='1F4E78')
NORM=Font(name=FONT,size=10)
thin=Side(style='thin',color='D9D9D9')
BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
CEN=Alignment(horizontal='center',vertical='center')
LEFT=Alignment(horizontal='left',vertical='center',wrap_text=True)
QFILL={'Q1':'C6E0B4','Q2':'FFE699','Q3':'BDD7EE','Q4':'F8CBAD'}

wb=Workbook()

def style_header(ws,row,ncol):
    for c in range(1,ncol+1):
        cell=ws.cell(row=row,column=c)
        cell.font=HDR; cell.fill=HDRFILL; cell.alignment=CEN; cell.border=BORD

def put(ws,r,c,v,font=None,al=None,fill=None,border=True,fmt=None):
    cell=ws.cell(row=r,column=c,value=v)
    cell.font=font or NORM
    if al: cell.alignment=al
    else: cell.alignment=Alignment(vertical='center')
    if fill: cell.fill=fill
    if border: cell.border=BORD
    if fmt: cell.number_format=fmt
    return cell

# ---------- Sheet 1: サマリー ----------
ws=wb.active; ws.title='サマリー'
ws.sheet_view.showGridLines=False
put(ws,1,1,'エクサベースAI 利用実態分析（2026年7月）',TITLE,border=False)
put(ws,2,1,'対象期間: 2026-07-01〜2026-07-31 ／ 全169ユーザー・2,587会話・19,548メッセージ',NORM,border=False)
t=agg['totals']
r=4
put(ws,r,1,'■ 全体サマリー',SUB,border=False); r+=1
summ=[('総ユーザー数',f"{t['users']} 名"),
      ('総会話数',f"{t['convs']:,} 件"),
      ('総メッセージ数（user+assistant）',f"{t['umsgs']+t['amsgs']:,} 件"),
      ('ユーザー発話数',f"{t['umsgs']:,} 件"),
      ('1会話あたり平均往復','3.4 往復（中央値2）'),
      ('Economy 利用（メッセージ）','16,063（82%）'),
      ('Standard 利用（メッセージ）','3,406（17%）'),
      ('Premium 利用（メッセージ）','8（0.04%）'),
      ('従量課金(Standard/Premium)会話数',f"404 件（うち過剰スペック候補 {agg['overspec_count']} 件）")]
for k,v in summ:
    put(ws,r,1,k,NORM); put(ws,r,2,v,Font(name=FONT,size=10,bold=True)); r+=1

r+=1
put(ws,r,1,'■ 主要インサイト',SUB,border=False); r+=1
insights=[
 '① 利用は二極化。全169名のうち高頻度・高リテラシー層(Q1)69名が全会話の78%を占有。一方で低頻度・低リテラシー層(Q4)67名は会話全体の1割未満で、シャドウAI利用が隠れている可能性が高い。',
 '② リテラシーは「利用頻度」とは相関する(r=0.60)が、「Standard/Premium(従量課金)の選択」とはほぼ無相関(r=0.08)。＝上位モデルは"必要だから"ではなく"なんとなく"選ばれている。',
 '③ 高度機能(Deep Research/RAG/ファイル出力等)の活用率はQ1でも僅か5%。リテラシー上位層でも高付加価値機能を使いこなせておらず、伸びしろが大きい。',
 '④ 従量課金404会話のうち81件(20%)は、単純用途(メール/校正/一般質問)×浅い往復×高度機能なしで、Economyで代替可能な過剰スペック候補。',
 '⑤ 用途はメール(16%)・画像編集(15%)・一般質問(12%)・営業リサーチ(10%)が中心。定型業務が多く、テンプレート化・共有で横展開しやすい。',
]
for s in insights:
    cell=put(ws,r,1,s,NORM,al=LEFT,border=False)
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
    ws.row_dimensions[r].height=45
    r+=1

r+=1
put(ws,r,1,'■ 分析手法・注記',SUB,border=False); r+=1
notes=[
 'リテラシー複合スコア(0-100)= 会話の深さ(平均往復)18% + プロンプト具体性(初回発話の長さ)18% + 高度ツール活用率20% + 用途の多様性16% + 構造化プロンプト率18% + 利用ツール種類数10%。各指標を169名内の百分位に変換し加重平均。',
 '利用頻度 = 月内のユーザー発話数。4象限は頻度・リテラシーそれぞれの中央値(頻度26/スコア62.5)で分割。',
 'チャット内容分類 = 各会話の初回発話をAIが読み、12カテゴリに分類(全2,587件)。',
 '過剰スペック候補 = Standard/Premiumモデルで、単純カテゴリ(C01/C02/C11/C12)かつ往復2回以下かつ高度ツール未使用の会話。',
]
for s in notes:
    cell=put(ws,r,1,s,Font(name=FONT,size=9,color='595959'),al=LEFT,border=False)
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
    ws.row_dimensions[r].height=42
    r+=1
ws.column_dimensions['A'].width=42
for col in 'BCDEF': ws.column_dimensions[col].width=16

# ---------- Sheet 2: ユーザー別 ----------
ws2=wb.create_sheet('ユーザー別')
ws2.sheet_view.showGridLines=False
heads=['ユーザー','利用頻度\n(発話数)','会話数','リテラシー\nスコア','4象限','平均往復','高度ツール\n活用率%','構造化\nプロンプト率%','用途\n多様性','ツール\n種類数','Std/Prem\n会話数','過剰スペック\n候補数','最多用途','Economy\n発話%','Std/Prem\n発話%']
for i,h in enumerate(heads,1): ws2.cell(row=1,column=i,value=h)
style_header(ws2,1,len(heads))
# sort by literacy desc
rows_sorted=sorted(rows,key=lambda x:(-x['literacy'],-x['user_msgs']))
rr=2
for u in rows_sorted:
    tiers=u['tiers']; tt=sum(tiers.values()) or 1
    ec=tiers.get('Economy',0)/tt*100
    sp=(tiers.get('Standard',0)+tiers.get('Premium',0))/tt*100
    topcat=max(u['cats'].items(),key=lambda x:x[1])[0] if u['cats'] else 'C12'
    q=u['quadrant'][:2]
    vals=[u['user'],u['user_msgs'],u['convs'],u['literacy'],u['quadrant'][3:],
          round(u['avg_turns'],1),round(u['adv_rate']*100,1),round(u['struct_rate']*100,1),
          u['cat_div'],u['distinct_tools'],u['std_prem_convs'],u['overspec_convs'],
          CAT[topcat],round(ec,0),round(sp,0)]
    for i,v in enumerate(vals,1):
        cell=put(ws2,rr,i,v,NORM,al=CEN if i not in(1,5,13) else LEFT)
        if i==5: cell.fill=PatternFill('solid',fgColor=QFILL[q])
        if i==4:
            cell.font=Font(name=FONT,size=10,bold=True)
    rr+=1
ws2.freeze_panes='A2'
ws2.auto_filter.ref=f'A1:{get_column_letter(len(heads))}{rr-1}'
widths=[26,10,8,10,20,8,10,11,8,8,10,11,26,10,10]
for i,w in enumerate(widths,1): ws2.column_dimensions[get_column_letter(i)].width=w
ws2.row_dimensions[1].height=42

# ---------- Sheet 3: パートナー別 ----------
ws3=wb.create_sheet('チャットパートナー別')
ws3.sheet_view.showGridLines=False
put(ws3,1,1,'チャットパートナー（モデル）別 利用状況',SUB,border=False)
heads=['チャットパートナー','モデル階層','会話数','ユーザー発話数','応答数','総メッセージ数']
for i,h in enumerate(heads,1): ws3.cell(row=3,column=i,value=h)
style_header(ws3,3,len(heads))
rr=4
TIERFILL={'Economy':'E2EFDA','Standard':'FFF2CC','Premium':'FCE4D6','Unknown':'FFFFFF'}
for p in agg['partners']:
    vals=[p['partner'],p['tier'],p['convs'],p['user_msgs'],p['asst_msgs'],p['total_msgs']]
    for i,v in enumerate(vals,1):
        cell=put(ws3,rr,i,v,NORM,al=LEFT if i==1 else CEN)
        if i==2: cell.fill=PatternFill('solid',fgColor=TIERFILL.get(p['tier'],'FFFFFF'))
    rr+=1
# total row
put(ws3,rr,1,'合計',Font(name=FONT,bold=True),al=LEFT)
put(ws3,rr,3,f'=SUM(C4:C{rr-1})',Font(name=FONT,bold=True),al=CEN)
put(ws3,rr,4,f'=SUM(D4:D{rr-1})',Font(name=FONT,bold=True),al=CEN)
put(ws3,rr,5,f'=SUM(E4:E{rr-1})',Font(name=FONT,bold=True),al=CEN)
put(ws3,rr,6,f'=SUM(F4:F{rr-1})',Font(name=FONT,bold=True),al=CEN)
ws3.column_dimensions['A'].width=44
for c in 'BCDEF': ws3.column_dimensions[c].width=15
ws3.row_dimensions[3].height=20

# ---------- Sheet 4: ツール別 ----------
ws4=wb.create_sheet('利用ツール別')
ws4.sheet_view.showGridLines=False
put(ws4,1,1,'利用ツール別 利用会話数（複合利用は各ツールにカウント）',SUB,border=False)
heads=['利用ツール','利用会話数']
for i,h in enumerate(heads,1): ws4.cell(row=3,column=i,value=h)
style_header(ws4,3,len(heads))
rr=4
for tl in agg['tools']:
    put(ws4,rr,1,tl['tool'],NORM,al=LEFT)
    put(ws4,rr,2,tl['convs'],NORM,al=CEN)
    rr+=1
put(ws4,rr+1,1,'※ ツールを使った会話は全体552件（21%）。高度機能の活用余地が大きい。',Font(name=FONT,size=9,color='595959'),al=LEFT,border=False)
ws4.column_dimensions['A'].width=40
ws4.column_dimensions['B'].width=14

# ---------- Sheet 5: 分類別 ----------
ws5=wb.create_sheet('チャット内容分類別')
ws5.sheet_view.showGridLines=False
put(ws5,1,1,'チャット内容分類別 利用状況とモデル階層内訳',SUB,border=False)
heads=['コード','業務カテゴリ','会話数','構成比%','総メッセージ','Economy','Standard','Premium','Std/Prem\n比率%']
for i,h in enumerate(heads,1): ws5.cell(row=3,column=i,value=h)
style_header(ws5,3,len(heads))
rr=4
totc=sum(c['convs'] for c in agg['cats'])
for c in agg['cats']:
    sp=c['standard']+c['premium']
    vals=[c['code'],c['name'],c['convs'],round(c['convs']/totc*100,1),c['msgs'],
          c['economy'],c['standard'],c['premium'],round(sp/c['convs']*100,1) if c['convs'] else 0]
    for i,v in enumerate(vals,1):
        put(ws5,rr,i,v,NORM,al=LEFT if i==2 else CEN)
    rr+=1
ws5.column_dimensions['A'].width=8
ws5.column_dimensions['B'].width=30
for c in 'CDEFGHI': ws5.column_dimensions[c].width=12
ws5.row_dimensions[3].height=28
ws5.freeze_panes='A4'

# ---------- Sheet 6: 4象限 ----------
ws6=wb.create_sheet('4象限サマリー')
ws6.sheet_view.showGridLines=False
put(ws6,1,1,'利用頻度 × リテラシー 4象限サマリー',SUB,border=False)
heads=['象限','人数','会話数','発話数','平均リテラシー','平均往復','高度ツール活用率%','Std/Prem会話','過剰スペック候補','推奨アクション']
for i,h in enumerate(heads,1): ws6.cell(row=3,column=i,value=h)
style_header(ws6,3,len(heads))
ACTIONS={
'Q1_高頻度・高リテラシー':'横展開の担い手。社内アンバサダー化・ベストプラクティス共有・高度機能(Deep Research/RAG)への引き上げ。',
'Q2_高頻度・低リテラシー':'量は使うが質に課題。プロンプト研修・テンプレート提供で費用対効果を改善。トレーニング効果が最も高い層。',
'Q3_低頻度・高リテラシー':'使えるが機会が少ない。業務への組み込み・活用シーンの提示で利用頻度を引き上げ。',
'Q4_低頻度・低リテラシー':'未活用層。オンボーディング・成功体験の設計が必要。シャドウAI利用のリスク調査対象。',
}
rr=4
for q in agg['quads']:
    key=q['quadrant']; qf=key[:2]
    vals=[key,q['users'],q['convs'],q['msgs'],q['avg_lit'],q['avg_turns'],q['avg_advrate'],q['stdprem'],q['overspec'],ACTIONS[key]]
    for i,v in enumerate(vals,1):
        cell=put(ws6,rr,i,v,NORM,al=LEFT if i in(1,10) else CEN)
        if i==1: cell.fill=PatternFill('solid',fgColor=QFILL[qf]); cell.font=Font(name=FONT,bold=True,size=10)
        if i==10: cell.alignment=LEFT
    ws6.row_dimensions[rr].height=54
    rr+=1
ws6.column_dimensions['A'].width=22
for c in 'BCDEFGHI': ws6.column_dimensions[c].width=11
ws6.column_dimensions['J'].width=60

# ---------- Sheet 7: 過剰スペック候補 ----------
ws7=wb.create_sheet('過剰スペック候補')
ws7.sheet_view.showGridLines=False
put(ws7,1,1,'従量課金(Standard/Premium)の過剰スペック候補一覧',SUB,border=False)
put(ws7,2,1,'条件: Standard/Premiumモデル × 単純用途(メール/校正/ネーミング/一般質問) × 往復2回以下 × 高度ツール未使用 → Economyで代替可能',Font(name=FONT,size=9,color='595959'),border=False)
# rebuild full overspec list
ADV={'Deep Research (詳細)','文書検索 (RAG)','スライド作成 (PowerPoint)','ファイル出力 (Word、Excelなど)','ウェブリサーチ (高速)','国内記事検索','官公需案件検索','音声書き起こし','画像生成 (高性能, GPT-image)','画像生成 (高速, Nano Banana)','SharePointファイル検索','Teams検索','Outlookカレンダー予定確認','Box検索','ヘルプページ検索'}
SIMPLE={'C01','C02','C11','C12'}
osp=[]
for c in convs:
    if c['tier'] in ('Standard','Premium') and c['cat'] in SIMPLE and c['user_msgs']<=2 and not any(t in ADV for t in c['tools']):
        osp.append(c)
heads=['ユーザー','モデル','階層','用途カテゴリ','往復','初回発話(抜粋)']
for i,h in enumerate(heads,1): ws7.cell(row=4,column=i,value=h)
style_header(ws7,4,len(heads))
rr=5
for c in sorted(osp,key=lambda x:x['user']):
    vals=[c['user'],c['model'],c['tier'],CAT[c['cat']],c['user_msgs'],(c['first_user_msg'] or '')[:90].replace('\n',' ')]
    for i,v in enumerate(vals,1):
        put(ws7,rr,i,v,NORM,al=LEFT if i in(1,2,4,6) else CEN)
    rr+=1
ws7.column_dimensions['A'].width=24
ws7.column_dimensions['B'].width=26
ws7.column_dimensions['C'].width=10
ws7.column_dimensions['D'].width=22
ws7.column_dimensions['E'].width=7
ws7.column_dimensions['F'].width=70
ws7.freeze_panes='A5'

wb.save(str(OUTPUT/'エクサベースAI_利用分析_202607.xlsx'))
print('saved xlsx, overspec rows:',len(osp))
